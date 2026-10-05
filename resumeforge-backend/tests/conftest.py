"""Shared pytest fixtures.

**No test depends on the teammate's real artifacts.** A tiny
``TfidfVectorizer(ngram_range=(1, 2)) + LogisticRegression`` pipeline is fitted on
a handful of synthetic sentences and written to a temporary directory; the API
fixtures point their settings at those temporary files. Every document fixture
(PDF/DOCX/TXT) is generated in memory - no real resume is used anywhere.
"""

from __future__ import annotations

import io
import sys
from collections.abc import Callable, Iterator, Sequence
from pathlib import Path
from typing import Any

import pytest

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:  # pragma: no cover - import side effect
    sys.path.insert(0, str(BACKEND_ROOT))

import joblib  # noqa: E402
import sklearn  # noqa: E402
from docx import Document  # noqa: E402
from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sklearn.feature_extraction.text import TfidfVectorizer  # noqa: E402
from sklearn.linear_model import LogisticRegression  # noqa: E402
from sklearn.pipeline import Pipeline  # noqa: E402

from api.main import create_app  # noqa: E402
from api.settings import Settings  # noqa: E402

SYNTHETIC_CLASSES: tuple[str, ...] = ("Accounting", "Data Science", "Human Resources")

SYNTHETIC_CORPUS: dict[str, tuple[str, ...]] = {
    "Accounting": (
        "accountant managing the general ledger month end close and bank reconciliation",
        "accounts payable clerk processing vendor invoices payments and bookkeeping reports",
        "financial analyst preparing budgets variance analysis and audit support work",
    ),
    "Data Science": (
        "data scientist building machine learning models with python pandas and numpy",
        "machine learning engineer training tensorflow models for nlp text classification",
        "applied scientist running sql experiments and feature engineering for forecasting",
    ),
    "Human Resources": (
        "human resources generalist recruiting screening interviews and onboarding new hires",
        "talent acquisition specialist managing candidate pipelines and employer branding",
        "hr business partner handling employee relations policy and performance reviews",
    ),
}

SAMPLE_RESUME_TEXT: str = (
    "Data scientist with five years of experience. Built and evaluated machine learning "
    "models in Python with pandas, numpy and scikit-learn, deployed them on AWS with a "
    "SQL feature store, and presented forecasting results to business stakeholders."
)

OTHER_RESUME_TEXT: str = (
    "Human resources generalist with six years in talent acquisition. Screened applicants, "
    "ran interview loops, onboarded new hires and maintained the HR policy handbook."
)

_METADATA_FILENAME = "metadata.json"
_MODEL_FILENAME = "model.joblib"
_VECTORIZER_FILENAME = "vectorizer.joblib"
_CLASSIFIER_FILENAME = "classifier.joblib"
_ENCODER_FILENAME = "label_encoder.joblib"


# --------------------------------------------------------------------------- #
# Synthetic model
# --------------------------------------------------------------------------- #


@pytest.fixture(scope="session")
def synthetic_classes() -> tuple[str, ...]:
    """The three synthetic classes used by the fixture model."""
    return SYNTHETIC_CLASSES


@pytest.fixture(scope="session")
def synthetic_corpus() -> dict[str, tuple[str, ...]]:
    """The synthetic training sentences grouped by class."""
    return SYNTHETIC_CORPUS


@pytest.fixture(scope="session")
def synthetic_pipeline() -> Pipeline:
    """Fit a tiny three-class TF-IDF + Logistic Regression pipeline."""
    texts: list[str] = []
    labels: list[str] = []
    for label, sentences in SYNTHETIC_CORPUS.items():
        texts.extend(sentences)
        labels.extend([label] * len(sentences))
    vectorizer = TfidfVectorizer(ngram_range=(1, 2), lowercase=True)
    classifier = LogisticRegression(max_iter=1000, random_state=0)
    return Pipeline([("tfidf", vectorizer), ("clf", classifier)]).fit(texts, labels)


@pytest.fixture(scope="session")
def model_dir(
    tmp_path_factory: pytest.TempPathFactory, synthetic_pipeline: Pipeline
) -> Path:
    """Write ``model.joblib`` + ``metadata.json`` into a temporary directory."""
    directory = tmp_path_factory.mktemp("models")
    joblib.dump(synthetic_pipeline, directory / _MODEL_FILENAME)
    (directory / _METADATA_FILENAME).write_text(
        '{"model_name": "SyntheticFixturePipeline", "trained_date": "2026-01-01",'
        f' "sklearn_version": "{sklearn.__version__}", "ngram_range": [1, 2]}}',
        encoding="utf-8",
    )
    return directory


@pytest.fixture(scope="session")
def part_model_dir(
    tmp_path_factory: pytest.TempPathFactory, synthetic_pipeline: Pipeline
) -> Path:
    """Write the vectorizer and classifier as two separate artifacts."""
    directory = tmp_path_factory.mktemp("models_parts")
    joblib.dump(synthetic_pipeline.named_steps["tfidf"], directory / _VECTORIZER_FILENAME)
    joblib.dump(synthetic_pipeline.named_steps["clf"], directory / _CLASSIFIER_FILENAME)
    return directory


@pytest.fixture(scope="session")
def encoder_model_dir(
    tmp_path_factory: pytest.TempPathFactory, synthetic_pipeline: Pipeline
) -> Path:
    """Write artifacts whose classifier predicts integer indices (needs an encoder)."""
    from sklearn.preprocessing import LabelEncoder

    directory = tmp_path_factory.mktemp("models_encoder")
    vectorizer = synthetic_pipeline.named_steps["tfidf"]
    classifier = LogisticRegression(max_iter=1000, random_state=0)

    texts: list[str] = []
    labels: list[str] = []
    for label, sentences in SYNTHETIC_CORPUS.items():
        texts.extend(sentences)
        labels.extend([label] * len(sentences))
    features = vectorizer.transform(texts)
    encoder = LabelEncoder().fit(labels)
    classifier.fit(features, encoder.transform(labels))

    joblib.dump(vectorizer, directory / _VECTORIZER_FILENAME)
    joblib.dump(classifier, directory / _CLASSIFIER_FILENAME)
    joblib.dump(encoder, directory / _ENCODER_FILENAME)
    return directory


# --------------------------------------------------------------------------- #
# Synthetic documents
# --------------------------------------------------------------------------- #


def _pdf_escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def _build_pdf(pages: Sequence[Sequence[str]]) -> bytes:
    """Build a minimal, valid, text-only PDF from ``pages`` of lines."""
def _build_pdf(pages: Sequence[Sequence[str]]) -> bytes:
    """Build a minimal, valid, text-only PDF with one page per entry in ``pages``.

    Object layout: ``1`` catalog, ``2`` page tree, ``3..`` page objects, then the
    font, then one content stream per page. A page with no lines produces an
    empty text layer, which mimics a scanned/blank page.
    """
    page_count = max(1, len(pages))
    pages = list(pages) or [[]]
    first_page_object = 3
    font_object = first_page_object + page_count
    first_stream_object = font_object + 1

    kids = " ".join(f"{first_page_object + index} 0 R" for index in range(page_count))
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode("latin-1"),
    }
    for index in range(page_count):
        objects[first_page_object + index] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_object} 0 R >> >> "
            f"/Contents {first_stream_object + index} 0 R >>"
        ).encode("latin-1")
    objects[font_object] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    for index, page_lines in enumerate(pages):
        operations = ["BT", "/F1 12 Tf", "14 TL", "72 720 Td"]
        for line in page_lines:
            operations.append(f"({_pdf_escape(line)}) Tj")
            operations.append("T*")
        operations.append("ET")
        stream = "\n".join(operations).encode("latin-1")
        objects[first_stream_object + index] = (
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = []
    for number in sorted(objects):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + objects[number] + b"\nendobj\n"

    xref_offset = len(out)
    count = len(objects) + 1
    out += f"xref\n0 {count}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {count} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    return bytes(out)


@pytest.fixture(scope="session")
def pdf_factory() -> Callable[..., bytes]:
    """Return a factory building synthetic PDFs in memory."""
    return _build_pdf


@pytest.fixture(scope="session")
def sample_pdf(pdf_factory: Callable[..., bytes]) -> bytes:
    """A one-page PDF containing the sample data-science resume text."""
    return pdf_factory([SAMPLE_RESUME_TEXT.split(". ")[:4]])


@pytest.fixture(scope="session")
def docx_factory() -> Callable[..., bytes]:
    """Return a factory building synthetic DOCX documents in memory."""

    def _build(paragraphs: Sequence[str], table_rows: Sequence[Sequence[str]]) -> bytes:
        document = Document()
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        if table_rows:
            table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
            for row_index, row in enumerate(table_rows):
                for cell_index, value in enumerate(row):
                    table.rows[row_index].cells[cell_index].text = value
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    return _build


@pytest.fixture(scope="session")
def sample_docx(docx_factory: Callable[..., bytes]) -> bytes:
    """A DOCX with paragraphs **and** a skills table."""
    return docx_factory(
        ["Data Scientist Resume", "Five years building machine learning models in Python."],
        [["Skills", "Python, pandas, scikit-learn, SQL, AWS"]],
    )


@pytest.fixture(scope="session")
def encrypted_pdf_bytes() -> bytes | None:
    """A password-protected PDF, or ``None`` when the library cannot build one."""
    from pypdf import PdfWriter

    try:
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        writer.encrypt("synthetic-password")
        buffer = io.BytesIO()
        writer.write(buffer)
        return buffer.getvalue()
    except Exception:  # pragma: no cover - depends on the pypdf build
        return None


# --------------------------------------------------------------------------- #
# Application fixtures
# --------------------------------------------------------------------------- #


@pytest.fixture
def reports_dir(tmp_path: Path) -> Path:
    """An empty reports directory (no evaluation files)."""
    directory = tmp_path / "reports"
    directory.mkdir(parents=True, exist_ok=True)
    return directory


@pytest.fixture
def settings_factory(model_dir: Path) -> Callable[..., Settings]:
    """Build :class:`Settings` pointing at the synthetic artifacts."""

    def _factory(**overrides: Any) -> Settings:
        values: dict[str, Any] = {
            "MODEL_PATH": str(model_dir / _MODEL_FILENAME),
            "VECTORIZER_PATH": str(model_dir / _VECTORIZER_FILENAME),
            "CLASSIFIER_PATH": str(model_dir / _CLASSIFIER_FILENAME),
            "LABEL_ENCODER_PATH": str(model_dir / _ENCODER_FILENAME),
            "MODEL_METADATA_PATH": str(model_dir / _METADATA_FILENAME),
            "REPORTS_DIR": "reports",
            "RATE_LIMIT": "1000/minute",
            "MIN_TEXT_CHARS": 50,
            "MAX_TEXT_CHARS": 100_000,
            "MAX_UPLOAD_MB": 5,
            "TOP_K": 3,
            "CONFIDENCE_THRESHOLD": 0.5,
            "ALLOWED_ORIGINS": ["http://localhost:5173"],
            "LOG_LEVEL": "CRITICAL",
            "ENVIRONMENT": "test",
        }
        values.update(overrides)
        return Settings(**values)

    return _factory


@pytest.fixture
def settings(settings_factory: Callable[..., Settings], reports_dir: Path) -> Settings:
    """Default settings for the API tests."""
    return settings_factory(REPORTS_DIR=str(reports_dir))


@pytest.fixture
def app(settings: Settings) -> FastAPI:
    """A FastAPI application wired to the synthetic artifacts."""
    return create_app(settings)


@pytest.fixture
def client(app: FastAPI) -> Iterator[TestClient]:
    """A ``TestClient`` with the application lifespan executed."""
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture
def unloaded_settings(tmp_path: Path) -> Settings:
    """Settings pointing at a directory that contains no artifacts."""
    empty = tmp_path / "empty_models"
    empty.mkdir(parents=True, exist_ok=True)
    return Settings(
        MODEL_PATH=str(empty / "model.joblib"),
        VECTORIZER_PATH=str(empty / "vectorizer.joblib"),
        CLASSIFIER_PATH=str(empty / "classifier.joblib"),
        LABEL_ENCODER_PATH=str(empty / "label_encoder.joblib"),
        MODEL_METADATA_PATH=str(empty / "metadata.json"),
        REPORTS_DIR=str(tmp_path / "empty_reports"),
        RATE_LIMIT="1000/minute",
        LOG_LEVEL="CRITICAL",
        ENVIRONMENT="test",
    )


@pytest.fixture
def unloaded_client(unloaded_settings: Settings) -> Iterator[TestClient]:
    """A client for an application whose model failed to load (503 endpoints)."""
    with TestClient(create_app(unloaded_settings)) as test_client:
        yield test_client


@pytest.fixture(autouse=True)
def _clear_classifier_cache() -> Iterator[None]:
    """Keep the process-wide classifier cache out of the tests' way."""
    from src.predict import clear_classifier_cache

    clear_classifier_cache()
    yield
    clear_classifier_cache()


@pytest.fixture
def synthetic_env(monkeypatch: pytest.MonkeyPatch, model_dir: Path) -> None:
    """Point the *environment* at the synthetic artifacts (used by the CLI tests)."""
    monkeypatch.setenv("MODEL_PATH", str(model_dir / _MODEL_FILENAME))
    monkeypatch.setenv("VECTORIZER_PATH", str(model_dir / _VECTORIZER_FILENAME))
    monkeypatch.setenv("CLASSIFIER_PATH", str(model_dir / _CLASSIFIER_FILENAME))
    monkeypatch.setenv("LABEL_ENCODER_PATH", str(model_dir / _ENCODER_FILENAME))
    monkeypatch.setenv("MODEL_METADATA_PATH", str(model_dir / _METADATA_FILENAME))
    monkeypatch.setenv("MIN_TEXT_CHARS", "50")
    monkeypatch.setenv("LOG_LEVEL", "CRITICAL")
